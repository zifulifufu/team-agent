"""Document library: chunks documents and indexes them so members in a group chat can
search them through the library_search tool.

Only the extracted text is kept (the original file is not). Retrieval is local BM25
(Chinese is tokenized into bigrams) and calls no cloud endpoint. Supports txt / md /
csv / json / html / pdf (with a text layer) / docx; scanned PDFs need OCR and are not
supported yet.
"""

from __future__ import annotations

from . import i18n

import io
import json
import re
import threading
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse

import httpx

from .store import Store
from .textindex import BM25, chunk_text, join_chunks, tokenize

MAX_BYTES = 30 * 1024 * 1024
MAX_CHARS = 3_000_000
TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".log", ".json", ".yaml", ".yml", ".xml", ".ini",
            ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".c", ".cpp", ".h", ".sql", ".sh", ".srt", ".vtt"}


class LibraryError(Exception):
    pass


class _Text(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):  # type: ignore[no-untyped-def]
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        elif tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "section"):
            self.parts.append("\n")

    def handle_endtag(self, tag):  # type: ignore[no-untyped-def]
        if tag in ("script", "style", "noscript") and self.skip:
            self.skip -= 1

    def handle_data(self, data):  # type: ignore[no-untyped-def]
        if not self.skip:
            self.parts.append(data)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Returns (kind, text)."""
    ext = Path(filename).suffix.lower()
    if len(data) > MAX_BYTES:
        raise LibraryError(i18n.pick_now(f"File is too large (limit {MAX_BYTES // 1024 // 1024} MB)", f"文件太大(上限 {MAX_BYTES // 1024 // 1024} MB)"))
    if ext in TEXT_EXT or not ext:
        text = _decode(data)
        if ext == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=1)
            except ValueError:
                pass
        return (ext.lstrip(".") or "txt"), text
    if ext in (".html", ".htm"):
        p = _Text()
        p.feed(_decode(data))
        return "html", re.sub(r"\n\s*\n+", "\n\n", "".join(p.parts))
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise LibraryError(i18n.pick_now("Reading PDF files needs pypdf: pip install pypdf", "读取 PDF 需要安装 pypdf:pip install pypdf")) from None
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [(p.extract_text() or "") for p in reader.pages]
        except Exception as e:  # noqa: BLE001
            raise LibraryError(i18n.pick_now(f"Could not parse the PDF: {e}", f"PDF 解析失败:{e}")) from None
        text = "\n\n".join(pages).strip()
        if not text:
            raise LibraryError(i18n.pick_now("This PDF has no extractable text (it may be a scan; OCR is not supported yet)", "这个 PDF 没有可提取的文字(可能是扫描件,暂不支持 OCR)"))
        return "pdf", text
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            raise LibraryError(i18n.pick_now("Reading docx files needs python-docx: pip install python-docx", "读取 docx 需要安装 python-docx:pip install python-docx")) from None
        try:
            d = docx.Document(io.BytesIO(data))
        except Exception as e:  # noqa: BLE001
            raise LibraryError(i18n.pick_now(f"Could not parse the docx: {e}", f"docx 解析失败:{e}")) from None
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "docx", "\n".join(parts)
    raise LibraryError(i18n.pick_now(f"{ext} files are not supported yet (convert to txt / md / pdf / docx first)", f"暂不支持 {ext} 格式(可以先转成 txt / md / pdf / docx)"))


DIR_EXT = TEXT_EXT | {".html", ".htm", ".pdf", ".docx"}
MAX_DIR_FILES = 300
MAX_URL_BYTES = 5 * 1024 * 1024


def fetch_url(url: str, timeout: float = 15.0) -> tuple[str, str, bytes]:
    """Download a web page or document -> (final url, content-type, content). Only http(s) is
allowed, at most 3 redirects, 5MB maximum."""
    u = urlparse(url.strip())
    if u.scheme not in ("http", "https") or not u.netloc:
        raise LibraryError(i18n.pick_now("A link has to start with http:// or https://", "链接需要以 http:// 或 https:// 开头"))
    try:
        with httpx.Client(timeout=timeout, follow_redirects=True, max_redirects=3,
                          headers={"User-Agent": "Mozilla/5.0 TeamAgent"}) as c:
            with c.stream("GET", url.strip()) as r:
                r.raise_for_status()
                buf = bytearray()
                for part in r.iter_bytes():
                    buf += part
                    if len(buf) > MAX_URL_BYTES:
                        raise LibraryError(i18n.pick_now(f"The page is too large (over {MAX_URL_BYTES // 1024 // 1024} MB)", f"页面太大(超过 {MAX_URL_BYTES // 1024 // 1024} MB)"))
                return str(r.url), r.headers.get("content-type", "").split(";")[0].strip().lower(), bytes(buf)
    except httpx.HTTPStatusError as e:
        raise LibraryError(i18n.pick_now(f"The server answered {e.response.status_code}", f"对方返回了 {e.response.status_code}")) from None
    except httpx.HTTPError as e:
        raise LibraryError(i18n.pick_now(f"Could not open this link: {type(e).__name__}", f"打不开这个链接:{type(e).__name__}")) from None


class Library:
    def __init__(self, store: Store):
        self.store = store
        self._bm: BM25 | None = None
        self._chunks: list[dict] = []
        self._dirty = True
        # Retrieval runs in a thread pool (see toolhub / api_ext), so rebuilding the index must
# hold a lock:
        # otherwise two threads both see _dirty=True and each builds one, and a reader may get a
# _chunks and a _bm that are not from the same build.
        self._index_lock = threading.Lock()

    def invalidate(self) -> None:
        self._dirty = True

    def _ensure(self) -> None:
        if not self._dirty and self._bm is not None:
            return
        with self._index_lock:
            if not self._dirty and self._bm is not None:   # another thread already built it while we waited for the lock
                return
            chunks = self.store.all_chunks()
            bm = BM25([tokenize(c["title"] + " " + c["text"]) for c in chunks])
            self._chunks, self._bm, self._dirty = chunks, bm, False

    # ------------------------------------------------------------------ write
    def add_text(self, title: str, text: str, filename: str = "", kind: str = "note", size: int | None = None,
                 did: str | None = None) -> dict:
        text = text.strip()
        if not text:
            raise LibraryError(i18n.pick_now("There is no usable text content", "没有可用的文本内容"))
        if len(text) > MAX_CHARS:
            raise LibraryError(i18n.pick_now(f"The text is too long (limit {MAX_CHARS} characters); split it before importing", f"文本太长(上限 {MAX_CHARS // 10000} 万字),请拆分后再导入"))
        chunks = chunk_text(text)
        doc = self.store.add_doc(title.strip() or filename or i18n.pick_now("Untitled", "未命名"), filename, kind,
                                 size if size is not None else len(text.encode()), chunks, did)
        self.invalidate()
        return doc

    def add_file(self, filename: str, data: bytes, title: str | None = None) -> dict:
        kind, text = extract_text(filename, data)
        return self.add_text(title or Path(filename).stem, text, filename, kind, len(data))

    def add_url(self, url: str) -> dict:
        final, ctype, data = fetch_url(url)
        if ctype == "application/pdf" or final.lower().split("?")[0].endswith(".pdf"):
            kind, text = extract_text("x.pdf", data)
            title = Path(urlparse(final).path).stem or urlparse(final).netloc
        elif ctype in ("", "text/html", "application/xhtml+xml") or "html" in ctype:
            kind, text = extract_text("x.html", data)
            m = re.search(rb"<title[^>]*>(.*?)</title>", data[:200_000], re.S | re.I)
            title = re.sub(r"\s+", " ", _decode(m.group(1))).strip() if m else ""
            title = title or urlparse(final).netloc + urlparse(final).path
        elif ctype.startswith("text/") or ctype in ("application/json", "application/xml"):
            text = _decode(data)
            title = Path(urlparse(final).path).name or urlparse(final).netloc
        else:
            raise LibraryError(i18n.pick_now(f"This content type is not supported yet: {ctype or 'unknown'}", f"暂不支持这种内容类型:{ctype or '未知'}"))
        return self.add_text(title[:120], text, final, "link", len(data))

    def add_dir(self, path: str, recursive: bool = True) -> dict:
        """Bulk import the documents in a folder. Re-importing the same folder: unchanged files are
skipped, files whose size changed are replaced with the new version."""
        root = Path(path.strip()).expanduser()
        if not root.is_absolute() or not root.is_dir():
            raise LibraryError(i18n.pick_now("Give the full path of a folder that exists", "请填写一个存在的文件夹的完整路径"))
        root = root.resolve()
        by_name = {d["filename"]: d for d in self.store.list_docs() if d["filename"]}
        added: list[dict] = []
        skipped: list[dict] = []
        seen = 0
        for p in sorted(root.rglob("*") if recursive else root.glob("*")):
            if p.is_symlink() or not p.is_file() or any(part.startswith(".") for part in p.relative_to(root).parts):
                continue
            if p.suffix.lower() not in DIR_EXT:
                continue
            seen += 1
            if seen > MAX_DIR_FILES:
                skipped.append({"name": "…", "reason": i18n.pick_now(f"Too many files; only the first {MAX_DIR_FILES} were processed",  # i18n-keep: placeholder row name; U+2026 is not Chinese
                                       f"文件太多,只处理前 {MAX_DIR_FILES} 个")})
                break
            rel = str(p.relative_to(root))
            try:
                size = p.stat().st_size
                old = by_name.get(str(p))
                if old and old["size"] == size:
                    skipped.append({"name": rel, "reason": i18n.pick_now("Already up to date", "已是最新")})
                    continue
                if size > MAX_BYTES:
                    raise LibraryError(i18n.pick_now(f"File is too large (limit {MAX_BYTES // 1024 // 1024} MB)", f"文件太大(上限 {MAX_BYTES // 1024 // 1024} MB)"))
                kind, text = extract_text(p.name, p.read_bytes())
                if not text.strip():
                    raise LibraryError(i18n.pick_now("There is no usable text content", "没有可用的文本内容"))   # check before touching the old version: if the new one is empty, keep the old
                if old:   # changed files: replaced in place, keeping document id, title and enabled state, so
# documents already selected in a group are not lost
                    self.delete(old["id"])
                    doc = self.add_text(old["title"], text, str(p), kind, size, did=old["id"])
                    if not old["enabled"]:
                        doc = self.update(doc["id"], {"enabled": False}) or doc
                else:
                    doc = self.add_text(p.stem, text, str(p), kind, size)
                added.append(doc)
            except (LibraryError, OSError) as e:
                skipped.append({"name": rel, "reason": str(e)})
        return {"added": added, "skipped": skipped}

    def update(self, did: str, patch: dict) -> dict | None:
        doc = self.store.update_doc(did, patch)
        self.invalidate()
        return doc

    def delete(self, did: str) -> None:
        self.store.delete_doc(did)
        self.invalidate()

    # ------------------------------------------------------------------- read
    def search(self, query: str, top_k: int = 5, doc_ids: list[str] | None = None) -> list[dict]:
        self._ensure()
        assert self._bm is not None
        q = tokenize(query)
        if not q or not self._chunks:
            return []
        allowed = set(doc_ids) if doc_ids is not None else None
        scored = [(i, s) for i, s in self._bm.scores(q).items()
                  if allowed is None or self._chunks[i]["doc_id"] in allowed]
        scored.sort(key=lambda x: -x[1])
        out = []
        for i, s in scored[:top_k]:
            c = self._chunks[i]
            out.append({"doc_id": c["doc_id"], "title": c["title"], "idx": c["idx"],
                        "text": c["text"], "score": round(s, 3)})
        return out

    def read(self, did: str, start: int = 0, limit: int = 3000) -> dict:
        doc = self.store.get_doc(did)
        if not doc:
            raise LibraryError(i18n.pick_now("That document does not exist", "文档不存在"))
        text = join_chunks([c["text"] for c in self.store.doc_chunks(did)])
        return {"doc": doc, "start": start, "end": min(start + limit, len(text)),
                "total": len(text), "text": text[start:start + limit]}

    def find_by_title(self, key: str) -> dict | None:
        key = key.strip().lower()
        docs = self.store.list_docs()
        return next((d for d in docs if d["id"] == key or d["title"].lower() == key), None) or \
            next((d for d in docs if key and key in d["title"].lower()), None)

    def scope_ids(self, ext_library: dict) -> list[str] | None:
        """Group settings -> searchable document ids. None = no restriction (all enabled ones);
[] = none may be used."""
        mode = (ext_library or {}).get("mode", "all")
        if mode == "off":
            return []
        if mode == "selected":
            return list(ext_library.get("ids", []))
        return None
