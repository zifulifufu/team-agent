import pytest

from app import document_preview, docwrite


@pytest.mark.skipif(not document_preview.available(), reason="document renderer not installed")
async def test_group_document_preview_renders_real_pages_and_records_coverage(tmp_path):
    from docx import Document
    from PIL import Image

    doc = Document()
    doc.add_paragraph("中文测试文档 第一页")
    doc.add_page_break()
    doc.add_paragraph("第二页 检查分批渲染")
    doc.save(tmp_path / "test.docx")
    result = await document_preview.render(tmp_path, "test.docx", start_page=2, max_pages=1)
    assert result["page_count"] == 2 and result["rendered_pages"] == [2, 2]
    assert len(result["sha256"]) == 64
    pictures = [f for f in result["files"] if f["kind"] == "image"]
    assert len(pictures) == 1 and pictures[0]["name"].endswith("page-2.png")
    im = Image.open(tmp_path / pictures[0]["name"])
    assert min(im.size) > 500
    assert all((tmp_path / f["name"]).is_file() for f in result["files"])


@pytest.mark.skipif(not document_preview.available(), reason="document renderer not installed")
async def test_exported_report_has_continuous_actual_page_numbers(tmp_path):
    import re
    from pypdf import PdfReader

    body = "\n\n".join(f"Synthetic paragraph {i}. " + "Layout verification text. " * 18
                       for i in range(25))
    docwrite.write(tmp_path, "numbered.docx", "docx", "Synthetic pagination test", body)
    result = await document_preview.render(tmp_path, "numbered.docx", max_pages=1)
    pdf = next(f for f in result["files"] if f["name"].endswith(".pdf"))
    pages = PdfReader(tmp_path / pdf["name"]).pages
    assert len(pages) >= 2
    for number, page in enumerate(pages, 1):
        assert re.search(rf"\b{number}\s*/\s*{len(pages)}\b", page.extract_text())
    assert "Synthetic paragraph 24" in pages[-1].extract_text()


async def test_document_preview_rejects_paths_and_bad_page_ranges(tmp_path):
    for path in ("../other.docx", "/etc/hosts", "missing.docx"):
        with pytest.raises(Exception):
            await document_preview.render(tmp_path, path)
    docwrite.write(tmp_path, "test.docx", "docx", "test", "body")
    for start, count in ((0, 1), (1, 0), (1, 13)):
        with pytest.raises(ValueError, match="start_page"):
            await document_preview.render(tmp_path, "test.docx", start, count)
    (tmp_path / "document-previews").symlink_to(tmp_path.parent, target_is_directory=True)
    with pytest.raises(ValueError, match="symlink"):
        await document_preview.render(tmp_path, "test.docx")


async def test_preview_is_a_group_tool_not_a_claim_of_visual_approval(store, make_router, monkeypatch):
    from tests.conftest import FakeLLM
    from tests.test_collab import setup
    monkeypatch.setattr(document_preview, "available", lambda: True)
    orch, g = setup(store, make_router, FakeLLM(default="OK"))
    ctx = await orch.toolhub.context(store.get_group(g["id"]), store.list_agents()[0], connect=False)
    assert "render_document" in ctx.tools

    async def fake_render(workspace, path, start_page, max_pages):
        return {"source": path, "page_count": 1, "rendered_pages": [1, 1], "files": []}
    monkeypatch.setattr(document_preview, "render", fake_render)
    out = await orch.toolhub.call(ctx, "render_document", {"path": "test.docx"})
    assert out.ok and "review_picture" in out.text and '"page_count": 1' in out.text
